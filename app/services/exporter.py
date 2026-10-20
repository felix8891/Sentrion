from __future__ import annotations
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from .audit import audit
from ..db import utcnow
from ..security import mask_account
from ..ui_text import ui_text, translate_json_text

NAVY = '16365C'
WHITE = 'FFFFFF'
LIGHT_BLUE = 'EAF1F8'

def _safe_cell(cell, value):
    if value is None:
        cell.value = ''
    elif isinstance(value, str):
        # Force strings that could be interpreted as formulas to stay text in Excel.
        cell.value = value
        if value.startswith(('=', '+', '-', '@')):
            cell.data_type = 's'
    else:
        cell.value = value

def _format_sheet(ws):
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(name='Aptos', bold=True, color=WHITE)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 32
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 12
        for cell in list(col)[:150]:
            if cell.value is not None:
                max_len = max(max_len, min(44, len(str(cell.value))))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 38)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F7F9FC')

def build_audit_workbook(conn) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Invoices'
    invoice_headers = ['invoice_id','invoice_receive_date','invoice_issue_date','invoice_number','vendor_id','vendor_company_name','vendor_address','vendor_tax_id','purchase_order_number','payment_due_date','currency','bank_name','bank_account_number','bank_account_holder_name','invoice_subtotal','invoice_discount','invoice_tax','invoice_total_price','risk_level','risk_reasons','verification_status','pic_approval_status','pic_approval_by','pic_approval_time','payment_status','payment_time','audit_id','document_hash']
    ws.append(invoice_headers)
    rows = conn.execute('''SELECT i.*, v.vendor_code, p.id as payment_id, p.created_at as payment_created,
        (SELECT id FROM audit_events a WHERE a.invoice_id=i.id ORDER BY a.created_at DESC LIMIT 1) AS latest_audit_id,
        (SELECT u.name FROM users u WHERE u.id=i.verified_by) AS verifier_name,
        (SELECT u.name FROM users u JOIN audit_events a ON a.actor_id=u.id WHERE a.invoice_id=i.id AND a.event_type IN ('INVOICE_APPROVED','GREEN_APPROVAL') ORDER BY a.created_at DESC LIMIT 1) AS approval_name,
        (SELECT a.created_at FROM audit_events a WHERE a.invoice_id=i.id AND a.event_type IN ('INVOICE_APPROVED','GREEN_APPROVAL') ORDER BY a.created_at DESC LIMIT 1) AS approval_time
        FROM invoices i LEFT JOIN vendors v ON v.id=i.vendor_id LEFT JOIN payments p ON p.invoice_id=i.id AND p.status='SIMULATED_SUCCESS'
        ORDER BY i.invoice_receive_date DESC''').fetchall()
    for r in rows:
        values = [r['id'],r['invoice_receive_date'],r['invoice_issue_date'],r['invoice_number'],r['vendor_id'],r['vendor_company_name'],r['vendor_address'],r['vendor_tax_id'],r['purchase_order_number'],r['due_date'],r['currency'],r['bank_name'],mask_account(r['bank_account_number']),r['bank_account_holder_name'],r['subtotal'],r['discount'],r['tax'],r['total'],r['risk_level'],translate_json_text(r['risk_reasons_json']),r['verification_status'],r['approval_status'],r['approval_name'],r['approval_time'],r['payment_status'],r['payment_created'],r['latest_audit_id'],r['document_hash']]
        ws.append(values)
    ws2 = wb.create_sheet('Invoice Items')
    ws2.append(['invoice_id','currency','invoice_item_id','line_number','item_name','sku','quantity','unit','unit_price','item_discount','tax_rate','item_tax','line_total'])
    for r in conn.execute('SELECT ii.*, i.currency FROM invoice_items ii JOIN invoices i ON i.id=ii.invoice_id ORDER BY ii.invoice_id, COALESCE(ii.line_number, 2147483647), ii.id').fetchall():
        ws2.append([r['invoice_id'],r['currency'],r['id'],r['line_number'],r['item_name'],r['sku'],r['quantity'],r['unit'],r['unit_price'],r['item_discount'],r['tax_rate'],r['item_tax'],r['line_total']])
    ws3 = wb.create_sheet('Audit Events')
    ws3.append(['audit_event_id','invoice_id','event_timestamp','actor_id','actor_role','event_type','event_result','risk_level_at_event','event_reason','related_approval_id','related_document_hash'])
    for r in conn.execute('SELECT * FROM audit_events ORDER BY created_at, id').fetchall():
        ws3.append([r['id'],r['invoice_id'],r['created_at'],r['actor_id'],r['actor_role'],r['event_type'],r['event_result'],r['risk_level_at_event'],ui_text(r['event_reason']),r['related_approval_id'],r['related_document_hash']])
    for sheet in wb.worksheets:
        _format_sheet(sheet)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                    cell.data_type = 's'
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out
