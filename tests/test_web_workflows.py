from pathlib import Path
from io import BytesIO
import json
import re
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from app.db import get_db
from .conftest import login_as, csrf_from

ROOT = Path(__file__).resolve().parents[1]


def vendor_evidence_files():
    identity = b'%PDF-1.4\nSentrion synthetic identity evidence fixture\n%%EOF'
    bank = b'%PDF-1.4\nSentrion synthetic bank ownership evidence fixture\n%%EOF'
    return [
        ('identity_files', ('identity-registration.pdf', identity, 'application/pdf')),
        ('bank_files', ('bank-ownership-confirmation.pdf', bank, 'application/pdf')),
    ]

def test_health_and_login_page(client: TestClient):
    assert client.get('/health').json()['status'] == 'ok'
    r = client.get('/login')
    assert r.status_code == 200 and 'Sign in to Sentrion' in r.text

def test_dashboard_invoice_vendor_po_pages(client: TestClient):
    login_as(client)
    for path, phrase in [('/dashboard','Financial Overview'),('/invoices','Invoice List'),('/vendors','Vendor directory'),('/purchase-orders','Saved purchase orders'),('/approvals','Request list'),('/audits','System events')]:
        response=client.get(path)
        assert response.status_code == 200, path
        assert phrase in response.text, path

def test_user_facing_pages_use_english_labels(client: TestClient):
    login_as(client)
    checks = {
        '/dashboard': ['Financial Overview', 'Demo environment', 'Audit activity'],
        '/invoices': ['Incoming Invoices', 'Invoice List', 'All risk levels', 'All verification statuses'],
        '/vendors': ['Vendor directory', 'Verification status'],
        '/purchase-orders': ['Upload PO', 'Saved purchase orders'],
        '/approvals': ['Sensitive approvals', 'Request list'],
        '/audits': ['System events', 'Export to Excel'],
    }
    for path, expected in checks.items():
        response = client.get(path)
        assert response.status_code == 200, path
        for text in expected:
            assert text in response.text, (path, text)
    assert 'Masuk ke Sentrion' not in client.get('/login').text

def test_upload_green_invoice_and_process_queue(client: TestClient):
    login_as(client)
    csrf=csrf_from(client,'/invoices')
    path=ROOT/'examples'/'invoices'/'01_invoice_green.pdf'
    response=client.post('/invoices/upload',data={'csrf_token':csrf},files=[('files',(path.name,path.read_bytes(),'application/pdf'))],follow_redirects=False)
    assert response.status_code==303 and 'uploaded=1' in response.headers['location']
    with get_db() as conn:
        row=conn.execute("SELECT * FROM invoices WHERE invoice_number='INV-MVP-GREEN-0001' ORDER BY created_at DESC LIMIT 1").fetchone()
        assert row is not None
        assert row['processing_status']=='COMPLETED'
        assert row['extraction_status']=='COMPLETED'
        assert row['risk_level']=='GREEN'
        assert row['document_hash'] and len(row['document_hash'])==64
        assert conn.execute('SELECT COUNT(*) n FROM invoice_items WHERE invoice_id=?',(row['id'],)).fetchone()['n']==2
        assert conn.execute('SELECT COUNT(*) n FROM invoice_documents WHERE invoice_id=?',(row['id'],)).fetchone()['n']==1
        invoice_id = row['id']
        trace_event = conn.execute("SELECT * FROM audit_events WHERE invoice_id=? AND event_type='EXTRACTION_ENGINE_SELECTED' ORDER BY created_at DESC, rowid DESC LIMIT 1", (invoice_id,)).fetchone()
        assert trace_event is not None
        trace_details = json.loads(trace_event['details_json'])
        assert trace_details['extraction_source'] == 'local_ocr_regex_demo'
        assert trace_details['request_status'] == 'disabled'
        assert trace_details['request_attempted'] is False
    list_page = client.get('/invoices?q=INV-MVP-GREEN-0001')
    detail_page = client.get(f'/invoices/{invoice_id}')
    assert 'INV-MVP-GREEN-0001' in list_page.text and 'Northstar Food Supplies Co.' in list_page.text
    assert 'INV-MVP-GREEN-0001' in detail_page.text and 'Northstar Food Supplies Co.' in detail_page.text
    assert 'Local OCR / regex parser' in detail_page.text
    assert 'Configured model (not called)' in detail_page.text
    assert 'Request: Disabled' in detail_page.text
    assert '<th>Extraction</th>' in list_page.text and '<th>Processing</th>' in list_page.text

    # The existing Re-analyze action must actually re-run extraction for an uploaded
    # source document, while retaining the original upload and prior audit events.
    detail_csrf = csrf_from(client, f'/invoices/{invoice_id}')
    reanalyze = client.post(f'/invoices/{invoice_id}/reanalyze', data={'csrf_token': detail_csrf}, follow_redirects=False)
    assert reanalyze.status_code == 303
    after_reanalyze_page = client.get(f'/invoices/{invoice_id}')
    assert 'Re-analysis has been queued.' in after_reanalyze_page.text
    assert 'data-invoice-processing-status="COMPLETED"' in after_reanalyze_page.text
    with get_db() as conn:
        row = conn.execute('SELECT * FROM invoices WHERE id=?', (invoice_id,)).fetchone()
        assert row['invoice_number'] == 'INV-MVP-GREEN-0001'
        assert row['processing_status'] == 'COMPLETED'
        assert row['extraction_status'] == 'COMPLETED'
        event_types = [r['event_type'] for r in conn.execute('SELECT event_type FROM audit_events WHERE invoice_id=?', (invoice_id,)).fetchall()]
        assert event_types.count('INVOICE_RECEIVED') == 1
        assert 'INVOICE_REPROCESS_QUEUED' in event_types
        assert event_types.count('INVOICE_PROCESSING_STARTED') >= 2


def test_actual_supplier_pdf_upload_persists_extraction_and_status_consistently(client: TestClient):
    login_as(client)
    csrf = csrf_from(client, '/invoices')
    path = ROOT / 'tests' / 'fixtures' / 'SENTRION_DOC_000001.pdf'
    response = client.post(
        '/invoices/upload', data={'csrf_token': csrf},
        files=[('files', (path.name, path.read_bytes(), 'application/pdf'))],
        follow_redirects=False,
    )
    assert response.status_code == 303

    with get_db() as conn:
        row = conn.execute(
            "SELECT * FROM invoices WHERE invoice_number='26-SI-000001' ORDER BY created_at DESC LIMIT 1"
        ).fetchone()
        assert row is not None
        assert row['processing_status'] == 'COMPLETED'
        assert row['extraction_status'] == 'COMPLETED'
        assert row['invoice_issue_date'] == '2026-05-16'
        assert row['due_date'] == '2026-05-30'
        assert row['vendor_company_name'] == 'UD Pusaka Persada Prima'
        assert row['vendor_address'] == 'Jalan Merdeka No. 16, Batam 50241, Indonesia'
        assert row['vendor_tax_id'] == 'TR-ID-33658-0063'
        assert row['vendor_id'] is None  # Extraction is not the same thing as vendor verification.
        assert row['purchase_order_number'] == 'PUR/26/12604'
        assert row['currency'] == 'IDR'
        assert row['subtotal'] == 39_092_351
        assert row['discount'] == 1_249_110
        assert row['tax'] == 4_162_755
        assert row['total'] == 42_005_996
        assert row['bank_name'] is None and row['bank_account_number'] is None
        invoice_id = row['id']
        extracted = json.loads(row['extracted_json'])
        assert extracted['buyer_company_name'] == 'CV Delta Trading Raya'
        assert extracted['buyer_tax_id'] == 'TR-ID-31755-2032'
        assert extracted['taxable_amount'] == 37_843_241
        assert extracted['payment_terms'] == 'Net 14'
        assert extracted['payment_reference'] == 'REM/26/72180'
        items = conn.execute(
            'SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY line_number', (invoice_id,)
        ).fetchall()
        assert len(items) == 8
        assert [item['sku'] for item in items] == [
            'FLR-010', 'FLR-005', 'SGR-048', 'SEA-001', 'SGR-042', 'SEA-010', 'OIL-040', 'SYR-037'
        ]
        assert [item['line_number'] for item in items] == list(range(1, 9))
        assert [item['item_discount'] for item in items] == [150_293, 331_863, 0, 0, 27_689, 0, 707_161, 32_104]
        assert conn.execute("SELECT COUNT(*) n FROM audit_events WHERE invoice_id=? AND event_type='INVOICE_PROCESSING_COMPLETED'", (invoice_id,)).fetchone()['n'] == 1

    list_page = client.get('/invoices?q=26-SI-000001')
    detail_page = client.get(f'/invoices/{invoice_id}')
    assert list_page.status_code == 200 and detail_page.status_code == 200
    assert '26-SI-000001' in list_page.text and 'UD Pusaka Persada Prima' in list_page.text
    assert '26-SI-000001' in detail_page.text and 'UD Pusaka Persada Prima' in detail_page.text
    assert 'CV Delta Trading Raya' in detail_page.text
    assert 'PUR/26/12604' in detail_page.text
    assert 'REM/26/72180' in detail_page.text
    assert 'SKU: FLR-010' in detail_page.text
    assert 'IDR 42,005,996' in detail_page.text


def test_green_approval_only_runs_mock_payment(client: TestClient):
    login_as(client)
    csrf=csrf_from(client,'/invoices/inv-demo-0002')
    response=client.post('/invoices/inv-demo-0002/approve',data={'csrf_token':csrf,'reason':'Demo approval test'},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        inv=conn.execute('SELECT * FROM invoices WHERE id=?',('inv-demo-0002',)).fetchone()
        payment=conn.execute("SELECT * FROM payments WHERE invoice_id=? AND status='SIMULATED_SUCCESS'",('inv-demo-0002',)).fetchone()
        assert inv['risk_level']=='GREEN'
        assert inv['payment_status']=='SIMULATED_SUCCESS'
        assert payment is not None
        assert payment['simulated_reference'].startswith('SIM-')

def test_yellow_cannot_be_paid_directly(client: TestClient):
    login_as(client)
    csrf=csrf_from(client,'/invoices/inv-demo-0046')
    response=client.post('/invoices/inv-demo-0046/approve',data={'csrf_token':csrf,'reason':'Attempt to bypass YELLOW'},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        inv=conn.execute('SELECT payment_status FROM invoices WHERE id=?',('inv-demo-0046',)).fetchone()
        assert inv['payment_status']!='SIMULATED_SUCCESS'

def test_excel_export_contains_three_real_data_sheets(client: TestClient):
    login_as(client)
    response=client.get('/export/audit.xlsx')
    assert response.status_code==200
    wb=load_workbook(BytesIO(response.content),read_only=True,data_only=True)
    assert wb.sheetnames==['Invoices','Invoice Items','Audit Events']
    item_headers=[c.value for c in next(wb['Invoice Items'].iter_rows(min_row=1,max_row=1))]
    assert 'currency' in item_headers
    ws=wb['Invoices']
    headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
    assert 'document_hash' in headers and 'risk_level' in headers and 'invoice_total_price' in headers and 'currency' in headers
    assert ws.max_row>=51

def test_same_person_cannot_second_approve_bank_change(client: TestClient):
    login_as(client,'supervisor@sentrion.local','SupervisorDemo!2026')
    csrf=csrf_from(client,'/vendors/vendor-001')
    response=client.post('/vendors/vendor-001/bank-change/request',data={
        'csrf_token':csrf,'bank_name':'Demo Test Bank','account_number':'998877665544','account_holder':'Northstar Food Supplies Co.',
        'reason':'Dual-authorization workflow test','evidence':'Confirmed through the official number registered for the local test.',
    },follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute("SELECT id FROM approval_requests WHERE request_type='BANK_ACCOUNT_CHANGE' AND entity_id='vendor-001' AND status='PENDING_SECOND_APPROVER' ORDER BY created_at DESC LIMIT 1").fetchone()
        assert req
    # The supervisor who submitted the request must not approve the same request as second approver.
    csrf=csrf_from(client,'/approvals')
    response=client.post(f"/approvals/{req['id']}/decision",data={'csrf_token':csrf,'decision':'APPROVE','reason':'Test attempt'},follow_redirects=False)
    assert response.status_code==403

def test_two_distinct_personnel_can_approve_vendor_bank_change(client: TestClient):
    login_as(client,'admin@sentrion.local','AdminDemo!2026')
    csrf=csrf_from(client,'/vendors')
    response=client.post('/vendors',data={
        'csrf_token':csrf,'vendor_code':'TEST-DUAL-01','company_name':'Demo Dual-Control Vendor Ltd.',
        'address':'1 Test Street, Demo City, USA','tax_id':'DEMO-TAX-TEST-01','country':'USA',
        'contact_person':'Taylor Test','email':'vendor-test@example.com','phone':'+1 555 0100',
        'bank_name':'Initial Demo Bank','bank_country':'USA','account_holder_name':'Demo Dual-Control Vendor Ltd.',
        'account_number':'887766554433','swift_bic':'DEMOTESTXXX',
        'identity_evidence':'Legal registration checked through the official business registry.',
        'bank_evidence':'Account ownership checked through the supplier official contact on file.',
        'reason':'New supplier needed for an integration test.',
    },files=vendor_evidence_files(),follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute("SELECT * FROM approval_requests WHERE request_type='VENDOR_ONBOARDING' AND status='PENDING_FIRST_APPROVER' ORDER BY created_at DESC LIMIT 1").fetchone()
        assert req and req['requested_by']=='usr-admin'
        req_id=req['id']
        assert conn.execute("SELECT 1 FROM vendors WHERE vendor_code='TEST-DUAL-01'").fetchone() is None
    # Maker cannot approve their own onboarding request.
    response=client.post(f'/approvals/{req_id}/decision',data={'csrf_token':csrf_from(client,'/approvals'),'decision':'APPROVE','reason':'Maker self-approval test'},follow_redirects=False)
    assert response.status_code==403
    # First approval by finance does not create or activate vendor records yet.
    login_as(client,'finance@sentrion.local','FinanceDemo!2026')
    csrf=csrf_from(client,f'/approvals/{req_id}')
    response=client.post(f'/approvals/{req_id}/decision',data={'csrf_token':csrf,'decision':'APPROVE','reason':'Reviewed legal registration evidence','review_confirmed':'yes'},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute('SELECT * FROM approval_requests WHERE id=?',(req_id,)).fetchone()
        assert req['status']=='PENDING_SECOND_APPROVER'
        assert conn.execute("SELECT 1 FROM vendors WHERE vendor_code='TEST-DUAL-01'").fetchone() is None
    # The same first reviewer cannot provide the final approval.
    response=client.post(f'/approvals/{req_id}/decision',data={'csrf_token':csrf_from(client,'/approvals'),'decision':'APPROVE','reason':'Attempt to approve twice'},follow_redirects=False)
    assert response.status_code==403
    login_as(client,'supervisor@sentrion.local','SupervisorDemo!2026')
    csrf=csrf_from(client,f'/approvals/{req_id}')
    response=client.post(f'/approvals/{req_id}/decision',data={'csrf_token':csrf,'decision':'APPROVE','reason':'Final review of vendor identity and bank ownership evidence','review_confirmed':'yes'},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        vendor=conn.execute("SELECT * FROM vendors WHERE vendor_code='TEST-DUAL-01'").fetchone()
        accounts=conn.execute('SELECT * FROM vendor_bank_accounts WHERE vendor_id=?',(vendor['id'],)).fetchall()
        req=conn.execute('SELECT * FROM approval_requests WHERE id=?',(req_id,)).fetchone()
        records=conn.execute('SELECT user_id,decision FROM approval_records WHERE request_id=?',(req_id,)).fetchall()
        assert vendor['current_bank_account_number']=='887766554433'
        assert vendor['verification_status']=='VERIFIED'
        assert vendor['country']=='USA' and vendor['email']=='vendor-test@example.com'
        assert len(accounts)==1 and accounts[0]['status']=='VERIFIED' and accounts[0]['swift_bic']=='DEMOTESTXXX'
        assert req['status']=='APPROVED'
        assert req['requested_by']=='usr-admin' and req['approved_by']=='usr-supervisor'
        assert {r['user_id'] for r in records}=={'usr-finance','usr-supervisor'}
        assert {r['decision'] for r in records}=={'APPROVE_FIRST','APPROVE'}


def test_vendor_onboarding_rejection_does_not_create_vendor(client: TestClient):
    login_as(client,'finance@sentrion.local','FinanceDemo!2026')
    csrf=csrf_from(client,'/vendors')
    response=client.post('/vendors',data={
        'csrf_token':csrf,'vendor_code':'TEST-REJECT-01','company_name':'Vendor Rejected Ltd.',
        'bank_name':'Review Bank','account_holder_name':'Vendor Rejected Ltd.','account_number':'987654321012',
        'identity_evidence':'Registration evidence supplied for rejection test.',
        'bank_evidence':'Bank ownership evidence supplied for rejection test.',
        'reason':'Create a request that should be rejected.',
    },files=vendor_evidence_files(),follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute("SELECT id FROM approval_requests WHERE request_type='VENDOR_ONBOARDING' AND status='PENDING_FIRST_APPROVER' ORDER BY created_at DESC LIMIT 1").fetchone()
    login_as(client,'supervisor@sentrion.local','SupervisorDemo!2026')
    csrf=csrf_from(client,f"/approvals/{req['id']}")
    response=client.post(f"/approvals/{req['id']}/decision",data={'csrf_token':csrf,'decision':'REJECT','reason':'Independent verification evidence was insufficient.','review_confirmed':'yes'},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        assert conn.execute("SELECT 1 FROM vendors WHERE vendor_code='TEST-REJECT-01'").fetchone() is None
        status=conn.execute('SELECT status FROM approval_requests WHERE id=?',(req['id'],)).fetchone()['status']
        assert status=='REJECTED'


def test_vendor_profile_change_requires_two_approvers(client: TestClient):
    login_as(client,'finance@sentrion.local','FinanceDemo!2026')
    csrf=csrf_from(client,'/vendors/vendor-001')
    response=client.post('/vendors/vendor-001/profile-change/request',data={
        'csrf_token':csrf,'company_name':'Northstar Food Supplies Co.','address':'Updated address for approval test',
        'tax_id':'DEMO-TAX-001','country':'USA','contact_person':'Alex Supplier','email':'supplier@example.com','phone':'5550101',
        'reason':'Update contact and address details.','evidence':'Supplier contact details checked against independent official record.',
    },follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute("SELECT * FROM approval_requests WHERE request_type='VENDOR_PROFILE_CHANGE' AND entity_id='vendor-001' ORDER BY created_at DESC LIMIT 1").fetchone()
        assert req and req['status']=='PENDING_FIRST_APPROVER'
        assert conn.execute("SELECT address FROM vendors WHERE id='vendor-001'").fetchone()['address']!='Updated address for approval test'
    login_as(client,'supervisor@sentrion.local','SupervisorDemo!2026')
    csrf=csrf_from(client,f"/approvals/{req['id']}")
    response=client.post(f"/approvals/{req['id']}/decision",data={'csrf_token':csrf,'decision':'APPROVE','reason':'First profile review completed.','review_confirmed':'yes'},follow_redirects=False)
    assert response.status_code==303
    login_as(client,'admin@sentrion.local','AdminDemo!2026')
    csrf=csrf_from(client,f"/approvals/{req['id']}")
    response=client.post(f"/approvals/{req['id']}/decision",data={'csrf_token':csrf,'decision':'APPROVE','reason':'Final profile evidence reviewed and approved.','review_confirmed':'yes'},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        vendor=conn.execute("SELECT * FROM vendors WHERE id='vendor-001'").fetchone()
        assert vendor['address']=='Updated address for approval test'
        assert conn.execute("SELECT status FROM approval_requests WHERE id=?",(req['id'],)).fetchone()['status']=='APPROVED'

def test_red_po_exception_requires_two_people_and_is_bound_to_amounts(client: TestClient):
    login_as(client,'finance@sentrion.local','FinanceDemo!2026')
    csrf=csrf_from(client,'/invoices/inv-demo-0048')
    response=client.post('/invoices/inv-demo-0048/red-resolution/request',data={
        'csrf_token':csrf,'reason':'The amended PO has been reviewed','evidence':'An amended PO record was independently confirmed for this demo.',
    },follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute("SELECT * FROM approval_requests WHERE request_type='RED_RESOLUTION' AND entity_id='inv-demo-0048' ORDER BY created_at DESC LIMIT 1").fetchone()
        assert req and req['status']=='PENDING_SECOND_APPROVER'
        request_id=req['id']
    login_as(client,'supervisor@sentrion.local','SupervisorDemo!2026')
    csrf=csrf_from(client,'/approvals')
    response=client.post(f'/approvals/{request_id}/decision',data={'csrf_token':csrf,'decision':'APPROVE','reason':'Amendment evidence reviewed'},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute('SELECT * FROM approval_requests WHERE id=?',(request_id,)).fetchone()
        invoice=conn.execute('SELECT risk_level FROM invoices WHERE id=?',('inv-demo-0048',)).fetchone()
        assert req['status']=='APPROVED'
        # The approved exception only affects the exact invoice and PO amounts that were reviewed.
        assert invoice['risk_level']=='GREEN'
        context=__import__('json').loads(req['context_json'])
        assert context['resolved_finding_codes']==['PO_MISMATCH']
        assert context['invoice_total'] and context['po_total']

def test_po_upload_stores_source_and_hash(client: TestClient):
    login_as(client)
    csrf=csrf_from(client,'/purchase-orders')
    path=ROOT/'examples'/'purchase_orders'/'01_purchase_order_upload_example.pdf'
    response=client.post('/purchase-orders/upload',data={'csrf_token':csrf},files={'file':(path.name,path.read_bytes(),'application/pdf')},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        po=conn.execute("SELECT * FROM purchase_orders WHERE po_number='PO-UPLOAD-8888'").fetchone()
        assert po is not None
        assert po['currency'] == 'USD'
        assert po['document_hash'] and len(po['document_hash'])==64
        assert po['document_path']
        assert conn.execute('SELECT COUNT(*) n FROM purchase_order_items WHERE po_id=?',(po['id'],)).fetchone()['n']==2
        po_id=po['id']
    response=client.get(f'/po-documents/{po_id}')
    assert response.status_code==200 and response.content.startswith(b'%PDF-')

def test_red_invoice_cannot_be_paid_by_direct_post(client: TestClient):
    login_as(client)
    csrf=csrf_from(client,'/invoices/inv-demo-0047')
    response=client.post('/invoices/inv-demo-0047/approve',data={'csrf_token':csrf,'reason':'Attempt to approve RED'},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        inv=conn.execute('SELECT risk_level,payment_status FROM invoices WHERE id=?',('inv-demo-0047',)).fetchone()
        assert inv['risk_level']=='RED'
        assert inv['payment_status']!='SIMULATED_SUCCESS'
        assert conn.execute("SELECT COUNT(*) n FROM payments WHERE invoice_id=? AND status='SIMULATED_SUCCESS'",('inv-demo-0047',)).fetchone()['n']==0


def test_sidebar_is_viewport_pinned_on_desktop(client: TestClient):
    """The authenticated shell exposes a fixed-height desktop sidebar and independent nav scroll."""
    login_as(client)
    response = client.get('/dashboard')
    assert response.status_code == 200
    assert 'class="sidebar-user"' in response.text
    assert 'class="sidebar-security"' in response.text
    css = client.get('/static/app.css?v=0.2.5')
    assert css.status_code == 200
    assert 'position: sticky' in css.text
    assert 'height: 100dvh' in css.text
    assert '.side-nav {' in css.text and 'overflow-y: auto' in css.text
    # Keep the mobile nav layout from inheriting the desktop sticky height.
    assert '@media (max-width: 700px)' in css.text
    assert 'height: auto' in css.text


def test_invoice_filters_apply_risk_verification_and_search(client: TestClient):
    """Risk, verification status, and text search must be conjunctive and use their own fields."""
    login_as(client)
    ids = ['inv-demo-0001', 'inv-demo-0002', 'inv-demo-0003']
    with get_db() as conn:
        originals = conn.execute(
            f"SELECT id, invoice_number, vendor_company_name, risk_level, verification_status FROM invoices WHERE id IN ({','.join('?' for _ in ids)})",
            ids,
        ).fetchall()
        assert len(originals) == 3
        conn.execute("UPDATE invoices SET invoice_number='FILTER-RED-VERIFIED', vendor_company_name='CV Filter Target', risk_level='RED', verification_status='VERIFIED' WHERE id=?", (ids[0],))
        conn.execute("UPDATE invoices SET invoice_number='FILTER-YELLOW-VERIFIED', vendor_company_name='CV Filter Target Yellow', risk_level='YELLOW', verification_status='VERIFIED' WHERE id=?", (ids[1],))
        conn.execute("UPDATE invoices SET invoice_number='FILTER-RED-PENDING', vendor_company_name='CV Filter Pending', risk_level='RED', verification_status='PENDING' WHERE id=?", (ids[2],))
    try:
        # Risk + verification + search must all apply at once.
        response = client.get('/invoices?q=FILTER&risk=RED&status=VERIFIED')
        assert response.status_code == 200
        assert '1 matching invoice.' in response.text
        assert 'FILTER-RED-VERIFIED' in response.text
        assert 'FILTER-YELLOW-VERIFIED' not in response.text
        assert 'FILTER-RED-PENDING' not in response.text
        assert 'risk-red">RED' in response.text

        # Search is case-insensitive and can match vendor company name.
        response = client.get('/invoices?q=filter%20target&risk=red&status=verified')
        assert response.status_code == 200
        assert 'FILTER-RED-VERIFIED' in response.text
        assert 'FILTER-YELLOW-VERIFIED' not in response.text
        assert '1 matching invoice.' in response.text

        # Verification status filters only the verification_status field, not processing/payment status.
        response = client.get('/invoices?q=FILTER&status=VERIFIED')
        assert response.status_code == 200
        assert '2 matching invoices.' in response.text
        assert 'FILTER-RED-VERIFIED' in response.text
        assert 'FILTER-YELLOW-VERIFIED' in response.text
        assert 'FILTER-RED-PENDING' not in response.text

        # A combination with no matches returns an empty result, not all invoices.
        response = client.get('/invoices?q=FILTER&risk=GREEN&status=VERIFIED')
        assert response.status_code == 200
        assert '0 matching invoices.' in response.text
        assert 'No matching invoices.' in response.text

        # The form clearly names the field being filtered, and Reset clears every query parameter.
        response = client.get('/invoices?risk=RED&status=VERIFIED')
        assert 'All verification statuses' in response.text
        assert 'href="/invoices" class="text-link" data-filter-reset>Reset</a>' in response.text
    finally:
        with get_db() as conn:
            for row in originals:
                conn.execute(
                    "UPDATE invoices SET invoice_number=?, vendor_company_name=?, risk_level=?, verification_status=? WHERE id=?",
                    (row['invoice_number'], row['vendor_company_name'], row['risk_level'], row['verification_status'], row['id']),
                )


def test_invoice_filters_update_without_visible_auto_hint_or_full_navigation(client: TestClient):
    """Live filters replace the results fragment without a full-document submit."""
    login_as(client)
    response = client.get('/invoices')
    assert response.status_code == 200
    assert 'data-live-invoice-filters' in response.text
    assert 'data-invoice-result-count' in response.text
    assert 'data-invoice-results' in response.text
    assert 'data-filter-reset' in response.text
    assert 'data-filter-error' in response.text
    assert 'Filter otomatis aktif' not in response.text
    assert 'data-filter-live-status' not in response.text
    # Extraction outcome and job state are separate, explicit columns.
    assert '<th>Verification</th>' in response.text
    assert '<th>Extraction</th>' in response.text
    assert '<th>Processing</th>' in response.text
    assert '<th>Approval</th>' in response.text

    js_response = client.get('/static/app.js?v=0.2.5')
    assert js_response.status_code == 200
    assert "form[data-live-invoice-filters]" in js_response.text
    assert "select.addEventListener('change'" in js_response.text
    assert "search.addEventListener('input'" in js_response.text
    assert 'setTimeout(() => {' in js_response.text
    assert 'fetch(url.toString()' in js_response.text
    assert "event.preventDefault()" in js_response.text
    assert "window.history.replaceState" in js_response.text
    assert "form.requestSubmit()" not in js_response.text
    assert "form.submit()" not in js_response.text

    css_response = client.get('/static/app.css?v=0.2.5')
    assert css_response.status_code == 200
    assert '.filter-error[hidden]' in css_response.text


def _make_invoice_pdf(text: str) -> bytes:
    import fitz
    document = fitz.open()
    page = document.new_page()
    page.insert_textbox(fitz.Rect(48, 48, 560, 740), text, fontsize=11, lineheight=1.4)
    data = document.tobytes()
    document.close()
    return data


def test_uploaded_invoice_with_missing_required_fields_is_partial_and_list_detail_match(client: TestClient):
    login_as(client)
    with get_db() as conn:
        vendor_count_before = conn.execute('SELECT COUNT(*) FROM vendors').fetchone()[0]
    csrf = csrf_from(client, '/invoices')
    content = _make_invoice_pdf("""Invoice No: PARTIAL-001
Invoice Date: 2026-10-17
Supplier: Northstar Food Supplies Co.
Currency: USD
Subtotal: USD 100.00
Tax: USD 8.00
Grand Total: USD 108.00
""")
    response = client.post('/invoices/upload', data={'csrf_token': csrf},
                           files=[('files', ('partial-invoice.pdf', content, 'application/pdf'))],
                           follow_redirects=False)
    assert response.status_code == 303
    with get_db() as conn:
        row = conn.execute("SELECT * FROM invoices WHERE invoice_number='PARTIAL-001' ORDER BY created_at DESC LIMIT 1").fetchone()
        assert row is not None
        assert row['processing_status'] == 'COMPLETED'
        assert row['extraction_status'] == 'PARTIAL'
        assert row['due_date'] is None
        assert row['discount'] is None
        invoice_id = row['id']
        event_types = [r['event_type'] for r in conn.execute('SELECT event_type FROM audit_events WHERE invoice_id=?', (invoice_id,)).fetchall()]
        assert 'INVOICE_RECEIVED' in event_types
        assert 'INVOICE_PROCESSING_STARTED' in event_types
        assert 'INVOICE_EXTRACTION_PARTIAL' in event_types
        assert conn.execute('SELECT COUNT(*) FROM vendors').fetchone()[0] == vendor_count_before
    list_page = client.get('/invoices?q=PARTIAL-001')
    detail_page = client.get(f'/invoices/{invoice_id}')
    assert 'PARTIAL-001' in list_page.text and 'Partial' in list_page.text
    assert 'PARTIAL-001' in detail_page.text and 'Partial' in detail_page.text
    assert 'JOB STATUS' in detail_page.text and 'Completed' in detail_page.text


def test_missing_invoice_number_stays_missing_in_list_and_detail(client: TestClient):
    login_as(client)
    csrf = csrf_from(client, '/invoices')
    content = _make_invoice_pdf("""Supplier: Example Unlisted Trading Co.
Currency: USD
Grand Total: USD 50.00
""")
    response = client.post('/invoices/upload', data={'csrf_token': csrf},
                           files=[('files', ('no-invoice-number.pdf', content, 'application/pdf'))],
                           follow_redirects=False)
    assert response.status_code == 303
    with get_db() as conn:
        row = conn.execute("SELECT * FROM invoices WHERE source_filename='no-invoice-number.pdf' ORDER BY created_at DESC LIMIT 1").fetchone()
        assert row is not None
        assert row['invoice_number'] is None
        assert row['extraction_status'] == 'PARTIAL'
        invoice_id = row['id']
    list_page = client.get('/invoices?q=Example%20Unlisted')
    detail_page = client.get(f'/invoices/{invoice_id}')
    assert 'Invoice number not extracted' in list_page.text
    assert 'Invoice number not extracted' in detail_page.text


def test_missing_uploaded_source_file_is_marked_failed_and_audited(client: TestClient):
    from uuid import uuid4
    from app.db import utcnow
    from app.services.processing import process_invoice

    login_as(client)
    invoice_id = 'inv-upload-missing-' + uuid4().hex[:12]
    now = utcnow()
    with get_db() as conn:
        user = conn.execute("SELECT id FROM users WHERE lower(email)=lower(?)", ('finance@sentrion.local',)).fetchone()
        conn.execute('''INSERT INTO invoices
            (id, invoice_receive_date, currency, extraction_status, processing_status, verification_status,
             approval_status, payment_status, risk_level, risk_reasons_json, source_filename, stored_path,
             document_hash, created_by, created_at, updated_at)
            VALUES (?, ?, 'USD', 'QUEUED', 'QUEUED', 'PENDING', 'NOT_REQUESTED', 'NOT_EXECUTED',
                    'YELLOW', '[]', ?, ?, ?, ?, ?, ?)''',
            (invoice_id, now, 'missing-source.pdf', '/path/that/does/not/exist/missing-source.pdf',
             '0' * 64, user['id'], now, now))

    process_invoice(invoice_id)
    with get_db() as conn:
        row = conn.execute('SELECT extraction_status, processing_status FROM invoices WHERE id=?', (invoice_id,)).fetchone()
        assert row['extraction_status'] == 'FAILED'
        assert row['processing_status'] == 'FAILED'
        events = [r['event_type'] for r in conn.execute('SELECT event_type FROM audit_events WHERE invoice_id=?', (invoice_id,)).fetchall()]
        assert 'INVOICE_PROCESSING_STARTED' in events
        assert 'INVOICE_PROCESSING_FAILED' in events


def test_uploaded_idr_invoice_currency_is_taken_from_document_not_default(client: TestClient):
    login_as(client)
    csrf = csrf_from(client, '/invoices')
    content = _make_invoice_pdf("""Invoice No: IDR-UPLOAD-001
Invoice Date: 2026-10-17
Due Date: 2026-11-16
Supplier: Harborview Packaging Ltd.
Currency: IDR
Subtotal: IDR 3,950
Tax: IDR 316
Grand Total: IDR 4,266
""")
    response = client.post('/invoices/upload', data={'csrf_token': csrf},
                           files=[('files', ('idr-invoice.pdf', content, 'application/pdf'))],
                           follow_redirects=False)
    assert response.status_code == 303
    with get_db() as conn:
        row = conn.execute("SELECT * FROM invoices WHERE invoice_number='IDR-UPLOAD-001' ORDER BY created_at DESC LIMIT 1").fetchone()
        assert row is not None
        assert row['currency'] == 'IDR'
        assert row['subtotal'] == 3950
        assert row['tax'] == 316
        assert row['total'] == 4266
        assert row['extraction_status'] == 'COMPLETED'


def test_vendor_onboarding_payload_tampering_is_blocked(client: TestClient):
    login_as(client,'finance@sentrion.local','FinanceDemo!2026')
    csrf=csrf_from(client,'/vendors')
    response=client.post('/vendors',data={
        'csrf_token':csrf,'vendor_code':'TEST-TAMPER-01','company_name':'Tamper Protection Ltd.',
        'bank_name':'Protection Bank','account_holder_name':'Tamper Protection Ltd.','account_number':'112233445566',
        'identity_evidence':'Legal supplier identity confirmed from an independent business registry.',
        'bank_evidence':'Bank account ownership independently confirmed through the official channel.',
        'reason':'Request created for payload-tampering regression test.',
    },files=vendor_evidence_files(),follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute("SELECT id,context_json FROM approval_requests WHERE request_type='VENDOR_ONBOARDING' AND status='PENDING_FIRST_APPROVER' ORDER BY created_at DESC, rowid DESC LIMIT 1").fetchone()
        context=json.loads(req['context_json'])
        context['vendor']['company_name']='Tampered Company Name Ltd.'
        conn.execute('UPDATE approval_requests SET context_json=? WHERE id=?',(json.dumps(context),req['id']))
    login_as(client,'supervisor@sentrion.local','SupervisorDemo!2026')
    response=client.post(f"/approvals/{req['id']}/decision",data={'csrf_token':csrf_from(client,'/approvals'),'decision':'APPROVE','reason':'Try approving modified payload'},follow_redirects=False)
    assert response.status_code==409
    with get_db() as conn:
        assert conn.execute("SELECT 1 FROM vendors WHERE vendor_code='TEST-TAMPER-01'").fetchone() is None
        assert conn.execute('SELECT status FROM approval_requests WHERE id=?',(req['id'],)).fetchone()['status']=='PENDING_FIRST_APPROVER'


def test_duplicate_vendor_tax_id_is_visible_and_requires_final_review_note(client: TestClient):
    login_as(client,'finance@sentrion.local','FinanceDemo!2026')
    csrf=csrf_from(client,'/vendors')
    response=client.post('/vendors',data={
        'csrf_token':csrf,'vendor_code':'TEST-DUP-TAX-01','company_name':'Separate Supplier with Similar Tax Data Ltd.',
        'tax_id':'DEMO-TAX-001','bank_name':'Separate Demo Bank','account_holder_name':'Separate Supplier Ltd.','account_number':'771122334455',
        'identity_evidence':'Legal identity record reviewed for duplicate-tax test.',
        'bank_evidence':'Bank ownership separately verified for duplicate-tax test.',
        'reason':'Request for duplicate tax ID warning regression test.',
    },files=vendor_evidence_files(),follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute("SELECT id,context_json FROM approval_requests WHERE request_type='VENDOR_ONBOARDING' AND status='PENDING_FIRST_APPROVER' AND context_json LIKE '%TEST-DUP-TAX-01%' ORDER BY created_at DESC, rowid DESC LIMIT 1").fetchone()
        context=json.loads(req['context_json'])
        assert 'TAX_ID_ALREADY_REGISTERED' in context['duplicate_flags']
    login_as(client,'supervisor@sentrion.local','SupervisorDemo!2026')
    response=client.get('/approvals')
    assert 'Tax Id Already Registered' in response.text
    csrf=csrf_from(client,f"/approvals/{req['id']}")
    response=client.post(f"/approvals/{req['id']}/decision",data={'csrf_token':csrf,'decision':'APPROVE','reason':'First reviewer looked at potential duplicate.','review_confirmed':'yes'},follow_redirects=False)
    assert response.status_code==303
    login_as(client,'admin@sentrion.local','AdminDemo!2026')
    csrf=csrf_from(client,f"/approvals/{req['id']}")
    response=client.post(f"/approvals/{req['id']}/decision",data={'csrf_token':csrf,'decision':'APPROVE','reason':'','review_confirmed':'yes'},follow_redirects=False)
    assert response.status_code==400
    with get_db() as conn:
        assert conn.execute("SELECT 1 FROM vendors WHERE vendor_code='TEST-DUP-TAX-01'").fetchone() is None
        assert conn.execute('SELECT status FROM approval_requests WHERE id=?',(req['id'],)).fetchone()['status']=='PENDING_SECOND_APPROVER'
    csrf=csrf_from(client,f"/approvals/{req['id']}")
    response=client.post(f"/approvals/{req['id']}/decision",data={'csrf_token':csrf,'decision':'APPROVE','reason':'Tax ID warning was examined; separate registration documents confirm distinct legal entity.','review_confirmed':'yes'},follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        assert conn.execute("SELECT 1 FROM vendors WHERE vendor_code='TEST-DUP-TAX-01'").fetchone()


def test_vendor_onboarding_evidence_uploads_and_reviewer_details(client: TestClient):
    login_as(client,'finance@sentrion.local','FinanceDemo!2026')
    csrf=csrf_from(client,'/vendors')
    response=client.post('/vendors',data={
        'csrf_token':csrf,'vendor_code':'TEST-EVIDENCE-01','company_name':'Evidence Review Supplier Ltd.',
        'address':'14 Document Way','tax_id':'EVIDENCE-TAX-001','country':'Indonesia',
        'contact_person':'Ari Evidence','email':'evidence@example.test','phone':'5551001',
        'bank_name':'Evidence Bank','bank_country':'Indonesia','account_holder_name':'Evidence Review Supplier Ltd.',
        'account_number':'123456789012','swift_bic':'EVDBIDJ1',
        'identity_evidence':'Legal registration independently checked against the official registry.',
        'bank_evidence':'Account ownership confirmed through the independently registered bank contact.',
        'reason':'Supplier required for the evidence-upload regression test.',
    },files=vendor_evidence_files(),follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute("SELECT * FROM approval_requests WHERE request_type='VENDOR_ONBOARDING' AND entity_id=(SELECT id FROM approval_requests WHERE request_type='VENDOR_ONBOARDING' AND context_json LIKE '%TEST-EVIDENCE-01%' ORDER BY created_at DESC LIMIT 1) ORDER BY created_at DESC LIMIT 1").fetchone()
        assert req is not None
        attachments=conn.execute('SELECT * FROM approval_attachments WHERE request_id=? ORDER BY category',(req['id'],)).fetchall()
        assert len(attachments)==2
        assert {a['category'] for a in attachments}=={'vendor_identity','bank_account'}
        request_id=req['id']
    approvals=client.get('/approvals')
    assert approvals.status_code==200 and f'/approvals/{request_id}' in approvals.text
    assert '2 evidence file(s)' in approvals.text
    login_as(client,'supervisor@sentrion.local','SupervisorDemo!2026')
    # Direct API submission cannot bypass reviewing the detail page, even with a forged checkbox value.
    direct=client.post(f'/approvals/{request_id}/decision',data={
        'csrf_token':csrf_from(client,'/approvals'),'decision':'APPROVE','reason':'Attempt without reviewing documents.','review_confirmed':'yes'
    },follow_redirects=False)
    assert direct.status_code==409
    detail=client.get(f'/approvals/{request_id}')
    assert detail.status_code==200
    for expected in ['Evidence Review Supplier Ltd.','EVIDENCE-TAX-001','Evidence Bank','123456789012',
                     'identity-registration.pdf','bank-ownership-confirmation.pdf','review_confirmed','I reviewed the complete vendor details']:
        assert expected in detail.text, expected
    with get_db() as conn:
        attachment=conn.execute("SELECT * FROM approval_attachments WHERE request_id=? AND category='vendor_identity'",(request_id,)).fetchone()
    served=client.get(f"/approvals/{request_id}/evidence/{attachment['id']}")
    assert served.status_code==200
    assert served.content.startswith(b'%PDF-1.4')
    assert served.headers['x-content-type-options']=='nosniff'
    # A request can be approved only through an explicit review acknowledgement.
    detail_csrf=csrf_from(client,f'/approvals/{request_id}')
    no_ack=client.post(f'/approvals/{request_id}/decision',data={
        'csrf_token':detail_csrf,'decision':'APPROVE','reason':'Reviewed supplier file details.'
    },follow_redirects=False)
    assert no_ack.status_code==400
    with get_db() as conn:
        assert conn.execute('SELECT status FROM approval_requests WHERE id=?',(request_id,)).fetchone()['status']=='PENDING_FIRST_APPROVER'


def test_vendor_onboarding_requires_identity_and_bank_files(client: TestClient):
    login_as(client,'finance@sentrion.local','FinanceDemo!2026')
    response=client.post('/vendors',data={
        'csrf_token':csrf_from(client,'/vendors'),'vendor_code':'TEST-NO-FILES','company_name':'No Files Supplier Ltd.',
        'bank_name':'Test Bank','account_holder_name':'No Files Supplier Ltd.','account_number':'998877665544',
        'identity_evidence':'Identity evidence narrative for the no-file validation test.',
        'bank_evidence':'Bank account evidence narrative for the no-file validation test.',
        'reason':'Test that documentary evidence is required.',
    },follow_redirects=False)
    assert response.status_code==400
    assert 'Upload at least one vendor identity document' in response.text
    with get_db() as conn:
        assert conn.execute("SELECT 1 FROM vendors WHERE vendor_code='TEST-NO-FILES'").fetchone() is None
        assert conn.execute("SELECT 1 FROM approval_requests WHERE request_type='VENDOR_ONBOARDING' AND context_json LIKE '%TEST-NO-FILES%' ").fetchone() is None


def test_tampered_or_modified_evidence_fails_integrity_check(client: TestClient):
    from app.config import UPLOAD_DIR
    login_as(client,'finance@sentrion.local','FinanceDemo!2026')
    response=client.post('/vendors',data={
        'csrf_token':csrf_from(client,'/vendors'),'vendor_code':'TEST-EVIDENCE-TAMPER','company_name':'Evidence Tamper Supplier Ltd.',
        'bank_name':'Tamper Bank','account_holder_name':'Evidence Tamper Supplier Ltd.','account_number':'776655443322',
        'identity_evidence':'Identity checked with the relevant corporate registry.',
        'bank_evidence':'Bank ownership checked using independent official channel.',
        'reason':'Testing immutable evidence attachment hashes.',
    },files=vendor_evidence_files(),follow_redirects=False)
    assert response.status_code==303
    with get_db() as conn:
        req=conn.execute("SELECT id FROM approval_requests WHERE request_type='VENDOR_ONBOARDING' AND context_json LIKE '%TEST-EVIDENCE-TAMPER%' ORDER BY created_at DESC LIMIT 1").fetchone()
        attachment=conn.execute("SELECT * FROM approval_attachments WHERE request_id=? AND category='vendor_identity'",(req['id'],)).fetchone()
    stored=UPLOAD_DIR/'vendor-evidence'/req['id']/attachment['stored_filename']
    stored.write_bytes(b'%PDF-1.4\\nTampered bytes\\n%%EOF')
    login_as(client,'supervisor@sentrion.local','SupervisorDemo!2026')
    detail=client.get(f"/approvals/{req['id']}")
    assert detail.status_code==200 and 'Missing or integrity check failed' in detail.text
    tampered_download=client.get(f"/approvals/{req['id']}/evidence/{attachment['id']}")
    assert tampered_download.status_code==409
    decision=client.post(f"/approvals/{req['id']}/decision",data={
        'csrf_token':csrf_from(client,'/approvals'),'decision':'APPROVE','reason':'Cannot approve tampered evidence.','review_confirmed':'yes'
    },follow_redirects=False)
    assert decision.status_code==409
    with get_db() as conn:
        assert conn.execute('SELECT status FROM approval_requests WHERE id=?',(req['id'],)).fetchone()['status']=='PENDING_FIRST_APPROVER'
